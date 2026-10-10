"""Build a job search tracker workbook (.xlsx) from the radar's database and config.

Tabs: Dashboard (formulas), Companies, Job Leads, Applications, Networking, Development
Plan and Search Config. The radar fills Companies and Job Leads; the other tabs start
empty, or with your own rows from <workspace>/tracker_seed.yaml:

    applications: [[ "APP-2024-01", "Company", "Role", ...], ...]   # one list per row, column order below
    networking: [...]
    development_plan: [...]
    search_config: [["Key", "Value", "Note"], ...]                  # extra rows

Upload the file to Google Drive and open it as a Google Sheet to get the layout the
optional sheet sync (radar/sheets.py) expects. Needs `pip install openpyxl`.

Usage: python tools/build_tracker.py [--workspace W] [--out tracker.xlsx] [--max-leads N] [--lean]
"""
from __future__ import annotations

import argparse
import datetime as dt
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "radar"))
import db  # noqa: E402
import workspace  # noqa: E402
from sources import ats_boards  # noqa: E402

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

LISTS = {
    "fit": ["High", "Medium", "Low", "Reject", "Pending"],
    "priority": ["P1", "P2", "P3"],
    "lead_status": ["Open", "Shortlisted", "Applied", "Interviewing", "Offer", "Not pursuing", "Closed"],
    "yesno": ["Yes", "No"],
    "stage": ["—", "Recruiter screen", "Hiring manager", "Take-home", "Technical 1", "Technical 2",
              "System design", "Final / Onsite", "Offer"],
    "outcome": ["—", "Pending", "Rejected", "Withdrew", "Ghosted", "Offer declined", "Offer accepted"],
    "app_status": ["Preparing", "Applied", "Interviewing", "Offer", "Rejected", "Withdrawn", "Closed"],
    "contact_type": ["Engineer / Team Member", "Hiring Manager", "In-house Recruiter", "Specialist Recruiter", "Other"],
    "channel": ["LinkedIn", "Internal messaging", "Email", "Event / In person", "Other"],
    "tier": ["A", "B", "C", "D"],
    "confidence": ["High", "Medium", "Low"],
    "dev_status": ["Not started", "In progress", "Done", "Paused"],
    "prio3": ["High", "Medium", "Low"],
}

APPLICATIONS = ["Application ID", "Company", "Job Title", "Job URL", "Date Applied", "Source Job ID", "Status",
                "Recruiter / Contact", "Recruiter Email", "CV Link", "Cover Letter Link", "Other Docs Link",
                "Application Folder Link", "Interview Dates", "Interview Notes", "Technical Topics", "Questions Missed",
                "Strengths / What Worked", "Outcome", "Compensation Discussed", "Follow-up Date",
                "Lessons / Reusable Insights", "Narrative Used", "Referral?"]
NETWORKING = ["Contact ID", "Job ID", "Company", "Job Title", "Contact Name", "Contact Role", "Contact Type", "Channel",
              "Profile / Contact Link", "Internal Contact?", "Date Identified", "Contacted?", "Date Contacted",
              "Response?", "Referral?", "Follow-up Date", "Relationship / Notes", "Next Action"]
DEV_PLAN = ["Area", "Priority", "Weekly Time", "Current Gap / Context", "Learning Goal", "Practice / Project",
            "Status", "Notes"]


def sheet(wb, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        c = ws.cell(1, i)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    return ws


def dropdown(ws, col_letter, key, n=2000):
    dv = DataValidation(type="list", formula1='"' + ",".join(LISTS[key]) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{n}")


def style_rows(ws, wrap_cols):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP if c.column_letter in wrap_cols else TOP


def build(out: Path, max_leads: int = 10000, lean: bool = False) -> tuple[int, int]:
    def cut(t, n=110):
        return ((t or "")[:n] + ("…" if t and len(t) > n else "")) if lean else (t or "")
    cfg = workspace.load_config()
    conf = (cfg.get("sources") or {}).get("ats_boards") or {}
    comps = ats_boards.companies(conf)
    seed_f = workspace.file("tracker_seed.yaml")
    seed = yaml.safe_load(seed_f.read_text(encoding="utf-8")) if seed_f.exists() else {}
    con = db.connect()
    today = dt.date.today().isoformat()
    wb = Workbook()

    ws = sheet(wb, "Dashboard", ["Metric", "Value", "How it is computed"], [42, 22, 70], first=True)
    L = "'Job Leads'"
    for r in [
        ("Open postings (Open + Shortlisted)", f'=COUNTIF({L}!J2:J5000,"Open")+COUNTIF({L}!J2:J5000,"Shortlisted")', "Job Leads, Status"),
        ("Open postings with a high fit", f'=COUNTIFS({L}!H2:H5000,"High",{L}!J2:J5000,"<>Closed",{L}!J2:J5000,"<>Not pursuing")', ""),
        ("Open postings with a medium fit", f'=COUNTIFS({L}!H2:H5000,"Medium",{L}!J2:J5000,"<>Closed",{L}!J2:J5000,"<>Not pursuing")', ""),
        ("Waiting for evaluation", f'=COUNTIF({L}!H2:H5000,"Pending")', "The radar evaluates a daily maximum"),
        ("Applications sent", '=COUNTIFS(Applications!E2:E1000,"<>",Applications!G2:G1000,"<>Preparing")', ""),
        ("Applications interviewing", '=COUNTIF(Applications!G2:G1000,"Interviewing")', ""),
        ("Offers", '=COUNTIF(Applications!G2:G1000,"Offer")+COUNTIF(Applications!S2:S1000,"Offer*")', ""),
        ("Interview rate", '=IF(B6<5,"Not enough data (<5)",(COUNTIF(Applications!G2:G1000,"Interviewing")+COUNTIF(Applications!G2:G1000,"Offer")+COUNTIF(Applications!S2:S1000,"Rejected"))/B6)',
         "Only computed with 5+ applications"),
        ("Contacts with a referral", '=COUNTIF(Networking!O2:O1000,"Yes")', ""),
        ("Active target companies", '=COUNTIF(Companies!M2:M500,"Yes")', ""),
        ("Snapshot date", today, ""),
    ]:
        ws.append(list(r))
    ws["B9"].number_format = "0%"
    ws.append([])
    ws.append(["By role family (open)", "", ""])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for fam in cfg.get("role_families") or []:
        if fam.get("active", True):
            ws.append([fam["name"], f'=COUNTIFS({L}!G2:G5000,"{fam["name"]}",{L}!J2:J5000,"<>Closed")', ""])
    style_rows(ws, {"C"})

    heads = ["Company", "Priority Tier", "Category", "Location Presence", "Target Role Families", "Why It Fits",
             "Typical Responsibilities", "Estimated Compensation", "Compensation Type", "Compensation Confidence",
             "Careers Page", "Current Notes", "Active Target?", "Hiring Process Notes", "ATS (radar)"]
    ws = sheet(wb, "Companies", heads, [20, 8, 22, 26, 26, 44, 34, 30, 16, 12, 30, 34, 9, 40, 20])
    for c in sorted(comps, key=lambda c: (c.get("tier") or "Z", c["company"])):
        ws.append([c["company"], c.get("tier"), c.get("category"), cut(c.get("presence"), 50),
                   ", ".join(c.get("families") or []), "" if lean else c.get("why"),
                   "" if lean else c.get("responsibilities"), cut(c.get("comp"), 60), c.get("comp_type"),
                   c.get("comp_confidence"), c.get("careers"), "" if lean else c.get("notes"),
                   "Yes" if c.get("active") else "No", "" if lean else c.get("hiring_notes", ""),
                   f"{c['ats'][0]}:{c['ats'][1]}" if c.get("ats") else "manual"])
    dropdown(ws, "B", "tier", 500); dropdown(ws, "M", "yesno", 500)
    style_rows(ws, set("CDEFGHLN"))
    ws.auto_filter.ref = f"A1:O{ws.max_row}"

    heads = ["Job ID", "First Seen", "Last Checked", "Company", "Job Title", "Location", "Role Family", "Fit",
             "Priority", "Status", "Applied?", "Date Applied", "Job URL", "Posting Date", "Estimated Compensation",
             "Compensation Type", "Responsibilities Summary", "Why It Fits", "Main Gaps", "CV Link",
             "Cover Letter Link", "Other Documents Link", "Application Folder Link", "Interview Stage", "Outcome", "Notes"]
    ws = sheet(wb, "Job Leads", heads, [16, 11, 11, 16, 36, 24, 20, 9, 8, 12, 9, 11, 30, 11, 22, 14, 40, 40, 36,
                                         14, 14, 14, 14, 14, 12, 36])
    ws.freeze_panes = "F2"
    leads = con.execute(
        """SELECT * FROM leads WHERE passes=1 AND status NOT IN ('Duplicate','Filtered')
           ORDER BY CASE fit WHEN 'High' THEN 0 WHEN 'Medium' THEN 1 WHEN 'Low' THEN 2 WHEN 'Reject' THEN 4 ELSE 3 END,
                    prescore DESC, company""").fetchall()
    if lean:
        leads = [r for r in leads if r["fit"] in ("High", "Medium")]
    for i, r in enumerate(leads[:max_leads], 1):
        status = "Closed" if r["status"] == "Closed" else ("Not pursuing" if r["fit"] == "Reject" else "Open")
        ws.append([f"JL-{i:04d}", r["first_seen"], r["last_checked"], r["company"], r["title"], (r["location"] or "")[:80],
                   r["family"], r["fit"] or "Pending", r["priority"] or "", status, "No", "", r["url"], r["posted"],
                   r["comp_estimate"] or r["comp"] or "", r["comp_type"] or "", cut(r["summary"]), cut(r["why_fits"], 140),
                   cut(r["main_gaps"]), "", "", "", "", "—", "—", cut(r["eval_notes"], 120) + f" [radar:{r['job_id']}]"])
    for col, key in (("H", "fit"), ("I", "priority"), ("J", "lead_status"), ("K", "yesno"), ("X", "stage"), ("Y", "outcome")):
        dropdown(ws, col, key, 5000)
    for val, color in (("High", "C6EFCE"), ("Medium", "FFEB9C"), ("Low", "EDEDED"), ("Reject", "F4CCCC")):
        ws.conditional_formatting.add("H2:H5000", CellIsRule(operator="equal", formula=[f'"{val}"'],
                                                             fill=PatternFill("solid", fgColor=color)))
    ws.conditional_formatting.add("A2:Z5000", FormulaRule(formula=['$J2="Closed"'], font=Font(color="999999")))
    style_rows(ws, set("EQRSZ"))
    ws.auto_filter.ref = f"A1:Z{max(ws.max_row, 2)}"

    ws = sheet(wb, "Applications", APPLICATIONS, [12, 14, 28, 26, 11, 12, 12, 20, 24, 14, 14, 14, 14, 18, 36, 24, 24,
                                                   30, 14, 16, 12, 44, 16, 9])
    for row in seed.get("applications") or []:
        ws.append(row)
    dropdown(ws, "G", "app_status", 1000); dropdown(ws, "S", "outcome", 1000); dropdown(ws, "X", "yesno", 1000)
    style_rows(ws, set("CHNOPQRV"))

    ws = sheet(wb, "Networking", NETWORKING, [10, 10, 14, 20, 20, 22, 18, 14, 26, 9, 11, 9, 11, 9, 9, 11, 40, 34])
    for row in seed.get("networking") or []:
        ws.append(row)
    for col, key in (("G", "contact_type"), ("H", "channel"), ("J", "yesno"), ("L", "yesno"), ("N", "yesno"), ("O", "yesno")):
        dropdown(ws, col, key, 1000)
    style_rows(ws, set("FQR"))

    ws = sheet(wb, "Development Plan", DEV_PLAN, [26, 9, 10, 42, 40, 44, 12, 30])
    for row in seed.get("development_plan") or []:
        ws.append(row)
    dropdown(ws, "B", "prio3", 200); dropdown(ws, "G", "dev_status", 200)
    style_rows(ws, set("ADEFH"))

    ws = sheet(wb, "Search Config", ["Key", "Value", "Note"], [34, 70, 40])
    fams = [f"{f['name']} ({'active' if f.get('active', True) else 'inactive'}, P{f.get('priority')})"
            for f in cfg.get("role_families") or []]
    cds = "; ".join(f"{c['company']} until {c['until']} ({c.get('note', '')})" for c in cfg.get("cooldowns") or [])
    for row in [("Role families", " · ".join(fams), ""), ("Cooldowns", cds, ""),
                ("Daily e-mail", (cfg.get("notify") or {}).get("to", ""), ""),
                ("Daily evaluation maximum", str((cfg.get("evaluation") or {}).get("max_per_day", 20)), "")]:
        ws.append(list(row))
    for row in seed.get("search_config") or []:
        ws.append(row)
    style_rows(ws, {"B", "C"})

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return len(leads[:max_leads]), len(leads)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--workspace")
    ap.add_argument("--out")
    ap.add_argument("--max-leads", type=int, default=10000)
    ap.add_argument("--lean", action="store_true", help="compact snapshot: only High/Medium leads, short texts")
    a = ap.parse_args()
    workspace.use(a.workspace)
    dest = Path(a.out) if a.out else workspace.root() / "out" / "Job_Search_Tracker.xlsx"
    n, tot = build(dest, a.max_leads, a.lean)
    print(f"leads in the workbook: {n}/{tot}; size: {dest.stat().st_size} bytes -> {dest}")
