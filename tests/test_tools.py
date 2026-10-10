"""tools/build_tracker.py builds the workbook from the database, config and optional seed."""
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

import db

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))
import build_tracker  # noqa: E402


def test_tracker_tabs_rows_and_seed(_workspace):
    shutil.copy(ROOT / "tests/config.yaml", _workspace / "config.yaml")
    (_workspace / "companies.yaml").write_text("- {company: Acme, tier: A, active: true, ats: [greenhouse, acme]}\n")
    (_workspace / "tracker_seed.yaml").write_text("applications:\n  - [APP-1, Acme, Backend Engineer]\n")
    con = db.connect()
    con.execute("INSERT INTO leads (job_id, company, title, passes, status, fit) VALUES ('gh:acme:1','Acme','Backend',1,'Evaluated','High')")
    con.commit()
    out = _workspace / "t.xlsx"
    assert build_tracker.build(out) == (1, 1)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Dashboard", "Companies", "Job Leads", "Applications", "Networking",
                             "Development Plan", "Search Config"]
    assert wb["Job Leads"]["Z2"].value.endswith("[radar:gh:acme:1]")
    assert wb["Applications"]["A2"].value == "APP-1"
